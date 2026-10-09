# -*- coding: utf-8 -*-
"""
core/report_make.py 单元测试（阶段 6 交付）

覆盖重点（对应方案阶段 6 的四条硬性要求）：
    1. TXT 分析报告：全文中文、固定四章结构
    2. 批量导出：数据表写为 xlsx（多 sheet），图片/日志/报告一并写入
    3. 命名规范：任务名_内容类型_YYYYMMDD_HHMMSS.扩展名
    4. 成果说明.txt：逐条说明每个文件是什么、怎么用
    另加：重复导出不覆盖、空数据边界、中文不乱码（BOM）、错误信息可操作。
"""

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "temp" / "mplconfig_test"))

from core import data_clean as dc  # noqa: E402
from core import data_parse as dp  # noqa: E402
from core import eco_index as ei  # noqa: E402
from core import plot_draw as pdraw  # noqa: E402
from core import report_make as rm  # noqa: E402

SAMPLES = ROOT / "tests" / "samples"


def _pipeline():
    """跑通导入 → 清洗 → 计算，返回各阶段结果。"""
    raw = dp.parse_files([SAMPLES / "sample_quadrat.csv"]).df
    cleaned = dc.clean_data(raw)
    index = ei.compute_indicators(cleaned.df, data_type=dp.TYPE_QUADRAT)
    return raw, cleaned, index


# ===========================================================================
# 一、命名规范
# ===========================================================================
class TestNaming:
    def test_base_name_format(self):
        name = rm.build_base_name("我的任务", "分析报告", timestamp="20240601_120000")
        assert name == "我的任务_分析报告_20240601_120000"

    def test_timestamp_is_added_automatically(self):
        name = rm.build_base_name("任务", "数据表")
        assert re.search(r"_\d{8}_\d{6}$", name), name

    def test_illegal_characters_removed(self):
        name = rm.build_base_name('任务/名:称*?"<>|', "报告", timestamp="20240601_120000")
        assert "/" not in name and ":" not in name and "*" not in name
        assert name.startswith("任务名称")

    def test_empty_task_name_falls_back(self):
        name = rm.build_base_name("", "报告", timestamp="20240601_120000")
        assert name.startswith(rm.DEFAULT_TASK_NAME)

    def test_unique_path_avoids_overwrite(self, tmp_path):
        first = rm.unique_path(tmp_path, "报告", ".txt")
        first.write_text("x", encoding="utf-8")
        second = rm.unique_path(tmp_path, "报告", ".txt")
        assert second != first
        assert second.name.startswith("报告_2")

    def test_unique_path_many_collisions(self, tmp_path):
        for _ in range(3):
            rm.unique_path(tmp_path, "报告", ".txt").write_text("x", encoding="utf-8")
        third = rm.unique_path(tmp_path, "报告", ".txt")
        assert third.name.startswith("报告_4"), third.name

    def test_suffix_without_dot(self, tmp_path):
        path = rm.unique_path(tmp_path, "报告", "txt")
        assert path.suffix == ".txt"


# ===========================================================================
# 二、分析报告（固定四章，全文中文）
# ===========================================================================
class TestReportText:
    def test_four_chapters_present(self):
        raw, cleaned, index = _pipeline()
        text = rm.build_report_text("测试任务", data_type=dp.TYPE_QUADRAT, raw=raw,
                                    clean=cleaned.df, row_index=index.row_table,
                                    summary_index=index.summary_table, clean_log=cleaned.log)
        for chapter in ("一、数据基本信息", "二、清洗统计摘要", "三、生态指标汇总",
                        "四、数据质量评价"):
            assert chapter in text, chapter

    def test_report_is_chinese_without_technical_noise(self):
        raw, cleaned, index = _pipeline()
        text = rm.build_report_text("测试任务", data_type=dp.TYPE_QUADRAT, raw=raw,
                                    clean=cleaned.df, clean_log=cleaned.log,
                                    row_index=index.row_table, summary_index=index.summary_table)
        for forbidden in ("Traceback", "DataFrame", "NaN", "inplace", "None", "dtype"):
            assert forbidden not in text, forbidden

    def test_chapter_one_has_counts(self):
        raw, cleaned, _index = _pipeline()
        text = rm.build_report_text("任务", raw=raw, clean=cleaned.df, clean_log=cleaned.log)
        assert "原始数据：共 38 行 × 5 列" in text
        assert "参与分析的数据：共 36 行 × 5 列" in text
        assert "样方号、物种、株数、胸径、盖度" in text

    def test_chapter_one_time_span_for_sensor(self):
        raw = dp.parse_files([SAMPLES / "sample_sensor.xlsx"]).df
        cleaned = dc.clean_data(raw, data_type=dp.TYPE_SENSOR)
        text = rm.build_report_text("传感器任务", data_type=dp.TYPE_SENSOR, clean=cleaned.df,
                                    clean_log=cleaned.log)
        assert "时间跨度：2024-05-01 00:00:00 至 2024-05-10 23:00:00" in text

    def test_chapter_two_numbers_match_cleaning_log(self):
        raw, cleaned, _index = _pipeline()
        text = rm.build_report_text("任务", raw=raw, clean=cleaned.df, clean_log=cleaned.log)
        assert "读取行数：38 行" in text
        assert "删除完全重复记录：2 行" in text
        assert "修正异常值：5 个" in text
        assert "清洗后行数：36 行" in text
        assert "突变检测阈值 3 倍 MAD" in text

    def test_chapter_three_has_metric_means(self):
        raw, cleaned, index = _pipeline()
        text = rm.build_report_text("任务", clean=cleaned.df, clean_log=cleaned.log,
                                    row_index=index.row_table,
                                    summary_index=index.summary_table)
        assert "群落多样性（按样方统计）" in text
        assert "H 香农指数：平均" in text
        assert "重要值 = (相对密度 + 相对频度 + 相对优势度) / 3" in text
        assert "优势物种" in text

    def test_chapter_three_explains_missing_dbh_rows(self):
        raw, cleaned, index = _pipeline()
        text = rm.build_report_text("任务", row_index=index.row_table,
                                    summary_index=index.summary_table)
        assert "没有胸径记录的行" in text
        assert "灌木" in text

    def test_chapter_four_reports_missing_rates(self):
        raw, cleaned, _index = _pipeline()
        text = rm.build_report_text("任务", clean=cleaned.df, clean_log=cleaned.log)
        assert "各字段缺失情况" in text
        assert "整体缺失率" in text
        assert "异常值处理：共修正 5 个" in text

    def test_chapter_four_does_not_claim_clean_when_blanks_exist(self):
        """有结构性空白时，不能说"未发现缺失值"——否则与上表自相矛盾。

        实测踩坑：示例数据里胸径有 50% 空白（灌木草本本来就没有），
        原报告却输出"本次清洗未发现异常值、缺失值与重复记录，数据本身是干净的"。
        """
        demo = dp.parse_files([ROOT / "static" / "demo" / "demo_quadrat.csv"]).df
        cleaned = dc.clean_data(demo)
        text = rm.build_report_text("任务", clean=cleaned.df, clean_log=cleaned.log)
        assert "胸径：空白 36 个（50.00%）" in text
        assert "数据本身是干净的" not in text
        assert "上表列出的空白属于该项本身没有观测" in text

    def test_chapter_four_says_clean_for_truly_clean_data(self):
        """确实没有任何空白与改动时才说"数据是干净的"。"""
        frame = pd.DataFrame({"样方号": ["Q1", "Q2"], "物种": ["甲", "乙"], "株数": [5, 8]})
        cleaned = dc.clean_data(frame)
        text = rm.build_report_text("任务", clean=cleaned.df, clean_log=cleaned.log)
        assert "数据本身是干净的" in text

    def test_report_without_data_still_valid(self):
        """没数据也要产出可用报告，而不是抛异常。"""
        text = rm.build_report_text("空任务")
        assert "一、数据基本信息" in text
        assert "本次报告未包含原始数据表" in text
        assert "报告结束" in text

    def test_report_includes_figure_list(self, tmp_path):
        raw, cleaned, index = _pipeline()
        chart = pdraw.draw_chart(index.summary_table,
                                 pdraw.recommend_spec(index.summary_table, "index_summary"))
        files = pdraw.save_figure(chart, tmp_path, formats=("png",))
        pdraw.close_figure(chart)
        text = rm.build_report_text("任务", figure_files=files)
        assert "本次导出的图表" in text
        assert files[0].name in text

    def test_parse_clean_stats_handles_garbage(self):
        """日志格式变化时解析失败不得抛异常，只是数字缺失。"""
        assert rm._parse_clean_stats("") == {}
        assert rm._parse_clean_stats("完全不是日志的内容") == {}
        stats = rm._parse_clean_stats("共读取 1,240 行数据。")
        assert stats.get("总行数") == "1,240"


# ===========================================================================
# 三、xlsx 写入
# ===========================================================================
class TestExcelWriting:
    def test_multi_sheet_and_chinese(self, tmp_path):
        raw, cleaned, index = _pipeline()
        target = tmp_path / "成果.xlsx"
        written = rm.write_tables_to_excel([
            ("清洗后数据", cleaned.df),
            ("行级指标", index.row_table),
            ("汇总指标", index.summary_table),
        ], target)
        assert written == target and target.exists()

        from openpyxl import load_workbook

        workbook = load_workbook(target)
        assert workbook.sheetnames == ["清洗后数据", "行级指标", "汇总指标"]
        sheet = workbook["汇总指标"]
        assert [cell.value for cell in sheet[1]] == list(index.summary_table.columns)
        assert sheet.freeze_panes == "A2"
        # 中文表头必须原样保留
        assert "H 香农指数" in [cell.value for cell in sheet[1]]

    def test_nan_written_as_blank(self, tmp_path):
        frame = pd.DataFrame({"物种": ["甲", "乙"], "胸径": [10.0, np.nan]})
        target = tmp_path / "含空值.xlsx"
        rm.write_tables_to_excel([("数据", frame)], target)
        from openpyxl import load_workbook

        sheet = load_workbook(target)["数据"]
        assert sheet.cell(row=2, column=2).value == 10.0
        assert sheet.cell(row=3, column=2).value is None

    def test_empty_tables_skipped_with_note(self, tmp_path):
        notes: list = []
        result = rm.write_tables_to_excel([("空表", pd.DataFrame())], tmp_path / "无.xlsx", notes)
        assert result is None
        assert notes and "没有可导出的数据表" in notes[0]
        assert not (tmp_path / "无.xlsx").exists()

    def test_sheet_name_sanitised(self):
        existing: list = []
        name = rm._safe_sheet_name("含[非法]:字符*的表/名", existing)
        assert not any(ch in name for ch in "[]:*?/\\")
        assert len(name) <= 31

    def test_sheet_name_deduplicated(self):
        existing = ["数据"]
        assert rm._safe_sheet_name("数据", existing) == "数据_2"
        assert rm._safe_sheet_name("数据", ["数据", "数据_2"]) == "数据_3"

    def test_long_sheet_name_truncated(self):
        assert len(rm._safe_sheet_name("很长的表名" * 20, [])) <= 31

    def test_column_width_autofit(self, tmp_path):
        frame = pd.DataFrame({"很长的中文列名在这里": ["值"], "B": ["短"]})
        target = tmp_path / "列宽.xlsx"
        rm.write_tables_to_excel([("数据", frame)], target)
        from openpyxl import load_workbook

        sheet = load_workbook(target)["数据"]
        assert sheet.column_dimensions["A"].width > 10


# ===========================================================================
# 四、批量导出
# ===========================================================================
class TestExportAll:
    def _export(self, tmp_path, **kwargs):
        raw, cleaned, index = _pipeline()
        params = dict(raw=raw, clean=cleaned.df, row_index=index.row_table,
                      summary_index=index.summary_table, clean_log=cleaned.log,
                      data_type=dp.TYPE_QUADRAT)
        params.update(kwargs)
        return rm.export_all(tmp_path / "成果", task_name="样方调查", **params)

    def test_exports_expected_files(self, tmp_path):
        result = self._export(tmp_path)
        names = [path.name for path in result.files]
        assert any(name.endswith(".xlsx") for name in names)
        assert any("清洗日志" in name for name in names)
        assert any("分析报告" in name for name in names)
        assert any("成果说明" in name for name in names)
        for path in result.files:
            assert path.exists() and path.stat().st_size > 0

    def test_export_result_reports_items(self, tmp_path):
        result = self._export(tmp_path)
        assert rm.ITEM_CLEAN_TABLE in result.items
        assert rm.ITEM_REPORT in result.items
        assert result.count() == len(result.files)

    def test_figure_files_are_copied(self, tmp_path):
        raw, cleaned, index = _pipeline()
        chart = pdraw.draw_chart(index.summary_table,
                                 pdraw.recommend_spec(index.summary_table, "index_summary"))
        figures = pdraw.save_figure(chart, tmp_path / "图形", formats=("png", "svg"))
        pdraw.close_figure(chart)

        result = rm.export_all(tmp_path / "成果", task_name="样方调查", clean=cleaned.df,
                               clean_log=cleaned.log, figure_files=figures)
        copied = [path for path in result.files if path.suffix in (".png", ".svg")]
        assert len(copied) == 2
        for path in copied:
            assert path.stat().st_size == next(
                source.stat().st_size for source in figures if source.suffix == path.suffix)

    def test_selected_items_only(self, tmp_path):
        result = self._export(tmp_path, items=[rm.ITEM_REPORT])
        names = [path.name for path in result.files]
        assert any("分析报告" in name for name in names)
        assert not any(name.endswith(".xlsx") for name in names)
        assert not any("清洗日志" in name for name in names)

    def test_missing_data_is_reported_not_silent(self, tmp_path):
        """勾选了但没有内容的成果项，必须说明为什么跳过。"""
        result = rm.export_all(tmp_path / "成果", task_name="空任务",
                               items=[rm.ITEM_CLEAN_TABLE, rm.ITEM_FIGURES])
        assert any("没有可导出的数据表" in note for note in result.notes)
        assert any("还没有导出过图片" in note for note in result.notes)

    def test_repeated_export_never_overwrites(self, tmp_path):
        raw, cleaned, _index = _pipeline()
        first = rm.export_all(tmp_path / "成果", task_name="任务", clean=cleaned.df,
                              clean_log=cleaned.log, timestamp="20240601_120000")
        second = rm.export_all(tmp_path / "成果", task_name="任务", clean=cleaned.df,
                               clean_log=cleaned.log, timestamp="20240601_120000")
        first_names = {path.name for path in first.files}
        second_names = {path.name for path in second.files}
        assert not (first_names & second_names), first_names & second_names
        for path in first.files:
            assert path.exists()

    def test_directory_is_created(self, tmp_path):
        nested = tmp_path / "a" / "b" / "c"
        result = rm.export_all(nested, task_name="任务")
        assert nested.exists() and result.directory == nested

    def test_result_always_includes_readme(self, tmp_path):
        result = rm.export_all(tmp_path / "成果", task_name="任务")
        assert any("成果说明" in path.name for path in result.files)


# ===========================================================================
# 五、中文编码（不乱码）
# ===========================================================================
class TestEncoding:
    def test_text_files_have_bom(self, tmp_path):
        """txt 一律 UTF-8 带 BOM：Excel / 记事本打开中文都不乱码。"""
        raw, cleaned, _index = _pipeline()
        result = rm.export_all(tmp_path / "成果", task_name="中文任务名",
                               clean=cleaned.df, clean_log=cleaned.log)
        for path in result.files:
            if path.suffix == ".txt":
                assert path.read_bytes().startswith(b"\xef\xbb\xbf"), path.name

    def test_chinese_survives_round_trip(self, tmp_path):
        raw, cleaned, _index = _pipeline()
        result = rm.export_all(tmp_path / "成果", task_name="中文任务名",
                               clean=cleaned.df, clean_log=cleaned.log)
        report = next(path for path in result.files if "分析报告" in path.name)
        text = report.read_text(encoding="utf-8-sig")
        assert "数据分析报告" in text
        assert "中文任务名" in text

    def test_chinese_in_excel(self, tmp_path):
        raw, cleaned, _index = _pipeline()
        result = rm.export_all(tmp_path / "成果", task_name="任务", clean=cleaned.df)
        xlsx = next(path for path in result.files if path.suffix == ".xlsx")
        from openpyxl import load_workbook

        sheet = load_workbook(xlsx)["清洗后数据"]
        assert sheet.cell(row=2, column=1).value == "Q1"
        assert sheet.cell(row=2, column=2).value == "油松"


# ===========================================================================
# 六、成果说明
# ===========================================================================
class TestReadme:
    def test_readme_lists_each_exported_file(self, tmp_path):
        raw, cleaned, index = _pipeline()
        result = rm.export_all(tmp_path / "成果", task_name="任务", raw=raw, clean=cleaned.df,
                               row_index=index.row_table, summary_index=index.summary_table,
                               clean_log=cleaned.log, data_type=dp.TYPE_QUADRAT)
        readme = next(path for path in result.files if "成果说明" in path.name)
        text = readme.read_text(encoding="utf-8-sig")
        for path in result.files:
            if path == readme:
                continue
            assert path.name in text, path.name

    def test_readme_has_purpose_and_howto(self, tmp_path):
        raw, cleaned, _index = _pipeline()
        result = rm.export_all(tmp_path / "成果", task_name="任务", clean=cleaned.df,
                               clean_log=cleaned.log)
        readme = next(path for path in result.files if "成果说明" in path.name)
        text = readme.read_text(encoding="utf-8-sig")
        assert "用途：" in text
        assert "怎么打开这些文件" in text
        assert ".xlsx" in text and ".svg" in text

    def test_readme_mentions_export_dir(self, tmp_path):
        result = rm.export_all(tmp_path / "成果", task_name="任务")
        readme = next(path for path in result.files if "成果说明" in path.name)
        text = readme.read_text(encoding="utf-8-sig")
        assert "导出位置：" in text
        assert str(result.directory) in text


# ===========================================================================
# 七、错误处理
# ===========================================================================
class TestErrors:
    def test_write_to_unwritable_location(self, tmp_path, monkeypatch):
        """目录不可创建时必须给出中文说明与建议，而不是原始异常。"""
        blocked = tmp_path / "只读" / "成果"

        def fake_mkdir(self, *args, **kwargs):
            raise OSError("模拟权限不足")

        monkeypatch.setattr(Path, "mkdir", fake_mkdir)
        with pytest.raises(rm.ReportError) as excinfo:
            rm.export_all(blocked, task_name="任务")
        message = str(excinfo.value)
        assert "无法创建输出文件夹" in message
        assert "怎么办" in message

    def test_report_error_message_is_chinese(self, tmp_path, monkeypatch):
        """导出的失败信息必须是中文、且带"怎么办"，不能把原始异常抛给用户。"""
        def fake_mkdir(self, *args, **kwargs):
            raise OSError("模拟磁盘已满")

        monkeypatch.setattr(Path, "mkdir", fake_mkdir)
        with pytest.raises(rm.ReportError) as excinfo:
            rm.export_all(tmp_path / "成果", task_name="任务")
        message = str(excinfo.value)
        assert any("\u4e00" <= ch <= "\u9fff" for ch in message), message
        for forbidden in ("Traceback", "OSError"):
            assert forbidden not in message


# ===========================================================================
# 八、图片文件名（避免双时间戳）
# ===========================================================================
class TestFigureNaming:
    def test_single_timestamp_even_when_time_differs(self, tmp_path):
        """实测踩坑：保存图片与导出成果常常不在同一秒。

        若只判断"源名里是否含本次时间戳"，会得到
        「图表名_200652_图表_200653.png」这种双时间戳文件名。
        本用例锁定"任何情况下文件名只含一个时间戳"。
        """
        figures_dir = tmp_path / "图"
        figures_dir.mkdir()
        source = figures_dir / "气温随时间变化_20240601_120000.png"
        source.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 100)

        # 导出时间戳与图片名里的时间戳不同
        result = rm.export_all(tmp_path / "成果", task_name="任务",
                               items=[rm.ITEM_FIGURES], figure_files=[source],
                               timestamp="20240602_093000")
        copied = [path for path in result.files if path.suffix == ".png"]
        assert len(copied) == 1
        name = copied[0].name
        assert name.count("2024") == 1, name
        assert name == "气温随时间变化_图表_20240602_093000.png", name

    def test_same_timestamp_also_single(self, tmp_path):
        figures_dir = tmp_path / "图"
        figures_dir.mkdir()
        source = figures_dir / "图表名_20240601_120000.png"
        source.write_bytes(b"x" * 50)
        result = rm.export_all(tmp_path / "成果", items=[rm.ITEM_FIGURES],
                               figure_files=[source], timestamp="20240601_120000")
        copied = [path for path in result.files if path.suffix == ".png"]
        assert copied[0].name == "图表名_图表_20240601_120000.png", copied[0].name

    def test_svg_keeps_extension(self, tmp_path):
        figures_dir = tmp_path / "图"
        figures_dir.mkdir()
        source = figures_dir / "图_20240601_120000.svg"
        source.write_text("<svg/>", encoding="utf-8")
        result = rm.export_all(tmp_path / "成果", items=[rm.ITEM_FIGURES],
                               figure_files=[source], timestamp="20240602_000000")
        assert any(path.suffix == ".svg" for path in result.files)
        assert all(path.exists() for path in result.files)

    def test_source_without_timestamp(self, tmp_path):
        figures_dir = tmp_path / "图"
        figures_dir.mkdir()
        source = figures_dir / "手工命名图.png"
        source.write_bytes(b"x" * 50)
        result = rm.export_all(tmp_path / "成果", items=[rm.ITEM_FIGURES],
                               figure_files=[source], timestamp="20240602_000000")
        copied = [path for path in result.files if path.suffix == ".png"]
        assert copied[0].name == "手工命名图_图表_20240602_000000.png", copied[0].name

    def test_all_item_names_are_defined(self):
        for item in rm.ALL_ITEMS:
            assert item in rm.ITEM_PURPOSE
            assert rm.ITEM_PURPOSE[item]
